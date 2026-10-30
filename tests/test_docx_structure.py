from pathlib import Path

import pytest
from docx import Document

from librarian.ingest.extractors import CompositeExtractor


@pytest.mark.asyncio
async def test_docx_preserves_headings_empty_cells_and_literal_pipes(tmp_path: Path) -> None:
    doc = Document()
    doc.add_heading("Results", level=2)
    table = doc.add_table(rows=3, cols=3)
    for cell, value in zip(table.rows[0].cells, ["Name", "Failures", "Successes"], strict=True):
        cell.text = value
    table.cell(1, 0).text = "Nora | Luis"
    table.cell(1, 2).text = "20"
    table.cell(2, 0).text = "Maya"
    table.cell(2, 1).text = "15"
    path = tmp_path / "results.docx"
    doc.save(str(path))
    text = await CompositeExtractor().extract(path)
    assert "## Results" in text
    assert "| Name | Failures | Successes |\n| --- | --- | --- |" in text
    assert r"| Nora \| Luis |  | 20 |" in text
    assert "| Maya | 15 |  |" in text
